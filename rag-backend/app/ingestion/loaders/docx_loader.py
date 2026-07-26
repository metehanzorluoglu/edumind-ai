import re
from pathlib import Path

import docx

from app.ingestion.errors import DocumentExtractionError
from app.ingestion.loaders.base import (
    ExtractedMetadata,
    ExtractionSource,
    LoadedDocument,
    PageContent,
)
from app.ingestion.metadata_extraction import is_plausible_title, split_author_string

_KEYWORD_SPLIT_PATTERN = re.compile(r"[;,]")


def load(path: Path) -> LoadedDocument:
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise DocumentExtractionError(f"Could not read DOCX file: {path.name}") from exc

    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    text = "\n".join(paragraphs).strip()

    if not text:
        raise DocumentExtractionError(f"No extractable text found in DOCX: {path.name}")

    core_properties = document.core_properties
    raw_title = (core_properties.title or "").strip()
    # Same reasoning as pdf_loader: Word's own "Save As" default can leave
    # the filename (or "Untitled") as the title core property.
    title = raw_title if raw_title and is_plausible_title(raw_title) else None

    author_raw = (core_properties.author or "").strip()
    authors = split_author_string(author_raw) if author_raw else []

    keywords_raw = (core_properties.keywords or "").strip()
    keywords = (
        [k.strip() for k in _KEYWORD_SPLIT_PATTERN.split(keywords_raw) if k.strip()]
        if keywords_raw
        else []
    )
    language = (core_properties.language or "").strip() or None

    sources: dict[str, ExtractionSource] = {}
    if title:
        sources["title"] = ExtractionSource.EMBEDDED_METADATA
    if authors:
        sources["authors"] = ExtractionSource.EMBEDDED_METADATA

    metadata = ExtractedMetadata(
        title=title, authors=authors, keywords=keywords, language=language, sources=sources
    )

    # python-docx has no layout engine, so real page boundaries (which depend on
    # page size, margins, and font metrics at render time) aren't available —
    # the whole document is one logical page rather than a fabricated count.
    return LoadedDocument(
        file_format="docx", pages=[PageContent(page_number=1, text=text)], metadata=metadata
    )
